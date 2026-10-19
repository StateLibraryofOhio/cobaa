import io

from django.urls import reverse
from openpyxl import load_workbook

from orders.models import Order


def test_export_creates_order_and_workbook(client, book):
    client.post(reverse('cart:cart_add', args=[book.id]), {'quantity': 2}, HTTP_REFERER='http://testserver/')
    r = client.post(reverse('pdf:books_to_excel'), {'inst_name': 'Test Library', 'amount_requested': '100'})
    assert r.status_code == 200
    order = Order.objects.get()
    assert order.items.get().quantity == 2
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws['A2'].value == book.title
    assert ws['C2'].value == 2


def test_export_invalid_form_400(client, book):
    r = client.post(reverse('pdf:books_to_excel'), {'inst_name': ''})
    assert r.status_code == 400
    assert Order.objects.count() == 0
